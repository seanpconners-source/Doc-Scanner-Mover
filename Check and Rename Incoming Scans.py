"""Rename completed checklist scans using their item number and PO/STOCK field."""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from datetime import datetime
import logging
import os
from pathlib import Path
import re
import shutil
import smtplib
import ssl
import tempfile
import time
import unicodedata
from email.message import EmailMessage

from openpyxl import load_workbook
from pdf2image import convert_from_path
from PIL import Image, ImageOps
from pypdf import PdfReader
import pytesseract


LOG = logging.getLogger("scan_sorter")
SUPPORTED_EXTENSIONS = {".pdf", ".png", ".jpg", ".jpeg"}
CHECKLIST_TITLE = re.compile(r"\bproduction\s+checkl[i1]st\b", re.I)
CHECKLIST_CELLS = (
    ("item-cell-wide", (.75, .13, .96, .165), 6),
    ("po-cell", (.17, .16, .57, .185), 6),
    ("item-cell-tight", (.81, .135, .93, .155), 7),
)
DEFAULT_INPUT = r"P:\AHAWKINS\Checklist Scans - AUTO SORT - DROP HERE"
DEFAULT_OUTPUT = r"P:\AHAWKINS\SORTED CHECKLISTS"
DEFAULT_ITEMS = r"C:\Users\sconners\Documents\Flavor List\FlavorList.xlsx"
DEFAULT_TESSERACT = r"C:\Users\Sconners\AppData\Local\Programs\Tesseract-OCR\tesseract.exe"
DEFAULT_POPPLER = r"C:\Users\Sconners\AppData\Local\Programs\Poppler\Library\bin"
DEFAULT_DEBUG = r"C:\Users\Sconners\Documents\Debug"


@dataclass(frozen=True)
class Config:
    input_folder: Path
    output_folder: Path
    debug_folder: Path
    excel_path: Path
    tesseract_path: str | None = None
    poppler_path: str | None = None
    language: str = "eng"
    dpi: int = 300
    ocr_timeout: int = 60
    poll_seconds: float = 5
    settle_seconds: float = 10
    retry_seconds: float = 60

    @classmethod
    def from_environment(cls):
        windows = os.name == "nt"
        return cls(
            input_folder=Path(os.environ.get("SCAN_INPUT_FOLDER", DEFAULT_INPUT)),
            output_folder=Path(os.environ.get("SCAN_OUTPUT_FOLDER", DEFAULT_OUTPUT)),
            debug_folder=Path(os.environ.get(
                "SCAN_DEBUG_FOLDER", DEFAULT_DEBUG if windows else
                str(Path(tempfile.gettempdir()) / "scan-sorter-debug"))),
            excel_path=Path(os.environ.get("SCAN_ITEM_LIST", DEFAULT_ITEMS)),
            tesseract_path=os.environ.get("SCAN_TESSERACT_PATH") or
                (DEFAULT_TESSERACT if windows and Path(DEFAULT_TESSERACT).is_file() else None),
            poppler_path=os.environ.get("SCAN_POPPLER_PATH") or
                (DEFAULT_POPPLER if windows and Path(DEFAULT_POPPLER).is_dir() else None),
        )


def normalize_item_number(value) -> str | None:
    """Excel often stores an eight-digit identifier as a number or '12345678.0'."""
    if value is None or isinstance(value, bool):
        return None
    value = str(value).strip()
    if re.fullmatch(r"\d{1,8}(?:\.0+)?", value):
        return value.split(".")[0].zfill(8)
    return None


def load_item_map(excel_path: Path | str) -> dict[str, str]:
    """Fail clearly rather than silently sorting every item as UNKNOWN."""
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    try:
        item_map = {}
        for row in workbook.worksheets[0].iter_rows(values_only=True):
            code = normalize_item_number(row[0]) if row else None
            if code is None:
                continue  # Header, blank line, or notes outside the item table.
            description = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
            if not description:
                raise ValueError(f"Item {code} has no description in {excel_path}")
            if code in item_map and item_map[code] != description:
                raise ValueError(f"Item {code} has conflicting descriptions in {excel_path}")
            item_map[code] = description
        if not item_map:
            raise ValueError(f"No eight-digit item numbers in the first column of {excel_path}")
        return item_map
    finally:
        workbook.close()


def sanitize_filename_component(name: str) -> str:
    name = re.sub(r'[\\/*?:"<>|\x00-\x1f]', "", name)
    return re.sub(r"\s+", " ", name).strip(" .")[:100].rstrip(" .")


# Corrections are permitted for item codes ONLY if the result exists in the workbook.
# No newlines between digits: otherwise unrelated numbers on separate lines merge.
ITEM_SYMBOLS = "0-9OoIl|SsBZJj"
ITEM_TOKEN = rf"[{ITEM_SYMBOLS}](?:[ \t-]*[{ITEM_SYMBOLS}]){{7}}"
ITEM_PATTERN = re.compile(rf"(?<![\w|])({ITEM_TOKEN})(?![\w|])")
ITEM_CORRECTIONS = str.maketrans({"O": "0", "o": "0", "I": "1", "i": "1", "l": "1", "L": "1",
                                "|": "1", "J": "1", "j": "1", "S": "5", "s": "5", "B": "8", "Z": "2"})
ITEM_LABEL = re.compile(
    r"\b(?:item|product|flavo[u]?r)\s*(?:no\.?|number|code|#)?\s*[:#=-]?\s*$", re.I)


def code_candidates(text: str, item_map: dict[str, str] | None = None):
    text = unicodedata.normalize("NFKC", text).replace("–", "-").replace("—", "-")
    candidates, labeled = set(), set()
    for match in ITEM_PATTERN.finditer(text):
        before, after = text[:match.start()], text[match.end():]
        if (re.search(r"(?<!\d)\d{1,7}[ \t-]+$", before) or
                re.match(r"[ \t-]+\d{1,7}(?!\d)", after)):
            continue  # Do not take eight digits out of a longer spaced number.
        token = re.sub(r"[ \t-]", "", match.group(1))
        if sum(c.isdigit() for c in token) < 4:
            continue
        if item_map is None and not token.isdigit():
            continue
        code = token.translate(ITEM_CORRECTIONS)
        if item_map is not None and code not in item_map:
            continue
        candidates.add(code)
        if ITEM_LABEL.search(text[max(0, match.start() - 50):match.start()]):
            labeled.add(code)
    return candidates, labeled


def extract_code_from_text(text, item_map=None):
    candidates, labeled = code_candidates(text, item_map)
    choices = labeled or candidates
    return next(iter(choices)) if len(choices) == 1 else None


PO_TOKEN = r"[0-9OoIl|SsBZ](?:[ \t-]*[0-9OoIl|SsBZ]){5}"
PO_PATTERN = re.compile(
    rf"\b(?:P[ .]*[O0][ .]*|purchase\s+order\s*)"
    rf"(?:no\.?\s*|number\s*)?[:#-]*\s*({PO_TOKEN})(?![\w|]|[ \t-]*\d)", re.I)


def po_candidates(text):
    text = unicodedata.normalize("NFKC", text)
    candidates = set()
    for match in PO_PATTERN.finditer(text):
        token = re.sub(r"[ \t-]", "", match.group(1))
        if sum(c.isdigit() for c in token) >= 4:
            candidates.add("PO" + token.translate(ITEM_CORRECTIONS))
    return candidates


def extract_po_from_text(text):
    candidates = po_candidates(text)
    if candidates:
        return next(iter(candidates)) if len(candidates) == 1 else None
    return "STOCK" if stock_field_found(text) else None


def stock_field_found(text, *, allow_standalone=True):
    # Full-page fallbacks must not mistake 'stock weight' or body instructions
    # mentioning stock for the purchase-order field.
    return bool((allow_standalone and re.search(r"^\s*STOCK\s*$", text, re.I | re.M)) or re.search(
        r"\b(?:P[ .]*[O0][ .]*|purchase\s+order\s*)"
        r"(?:no\.?\s*|number\s*)?[:#-]*\s*STOCK\b", text, re.I))


def select_fields(texts: dict[str, str], item_map: dict[str, str]):
    candidates, labeled, regional, exact_cells = set(), set(), set(), set()
    purchase_orders, cell_purchase_orders = set(), set()
    cell_stock = False
    is_checklist = any(CHECKLIST_TITLE.search(text) for text in texts.values())
    for name, text in texts.items():
        codes, labels = code_candidates(text, item_map)
        # Preserve a clearly labeled, exact item number even if the workbook is
        # out of date. Never infer an unknown item from arbitrary full-page dates.
        exact_codes, exact_labels = code_candidates(text)
        candidates.update(codes)
        labeled.update(labels)
        labeled.update(exact_labels)
        if name.startswith(("item-crop", "item-cell")):
            regional.update(codes)
        if name.startswith("item-cell"):
            exact_cells.update(exact_codes)
        purchase_orders.update(po_candidates(text))
        if name.startswith("po-cell"):
            cell_purchase_orders.update(po_candidates(text))
            cell_stock = cell_stock or stock_field_found(text)
    choices = labeled or regional or candidates or exact_cells
    code = next(iter(choices)) if len(choices) == 1 else None
    orders = cell_purchase_orders or purchase_orders
    po = next(iter(orders)) if len(orders) == 1 else None
    if cell_stock:
        po = None if cell_purchase_orders else "STOCK"
    elif not orders and any(stock_field_found(text, allow_standalone=not is_checklist)
                           for text in texts.values()):
        po = "STOCK"
    return code, po


def preprocess_image(image):
    # Preserve thin/faint characters. The old median filter and fixed threshold
    # removed their strokes, especially after the second JPEG/PDF conversion.
    return ImageOps.autocontrast(image.convert("L"))


def crop_code_area(image):
    w, h = image.size
    return image.crop((int(w * .6), 0, w, int(h * .25)))


def po_primary_crop(image):
    w, h = image.size
    return image.crop((0, 0, int(w * .5), int(h * .5)))


def checklist_cell(image, box):
    """Crop a value inside the checklist grid, excluding bordering rules."""
    w, h = image.size
    coordinates = tuple(int(value * (w if index % 2 == 0 else h))
                        for index, value in enumerate(box))
    with image.crop(coordinates) as cell:
        return ImageOps.expand(cell, border=20, fill="white")


def read_first_page(path: Path, config: Config) -> Image.Image:
    if path.suffix.lower() == ".pdf":
        pages = convert_from_path(str(path), dpi=config.dpi, first_page=1, last_page=1,
                                  poppler_path=config.poppler_path, thread_count=1,
                                  timeout=config.ocr_timeout)
        if not pages:
            raise ValueError("PDF has no readable first page")
        try:
            return pages[0].convert("RGB")
        finally:
            for page in pages:
                page.close()
    with Image.open(path) as original:
        oriented = ImageOps.exif_transpose(original)
        # Transparent backgrounds otherwise turn black and obscure the text.
        if "A" in oriented.getbands() or "transparency" in oriented.info:
            rgba = oriented.convert("RGBA")
            white = Image.new("RGBA", rgba.size, "white")
            white.alpha_composite(rgba)
            return white.convert("RGB")
        return oriented.convert("RGB")


def recognize_image(image, item_map, config, texts, debug_dir):
    grayscale = image.convert("L")
    contrast = preprocess_image(image)
    w, h = image.size
    # Tesseract's automatic segmentation can entirely discard text inside the
    # checklist grid. Detect the form, then read its value cells as text blocks.
    texts["full-page"] = pytesseract.image_to_string(
        grayscale, lang=config.language, config="--psm 3", timeout=config.ocr_timeout)
    is_checklist = bool(CHECKLIST_TITLE.search(texts["full-page"]))
    passes = []
    if is_checklist:
        passes.extend((name, checklist_cell(grayscale, box), mode)
                      for name, box, mode in CHECKLIST_CELLS)
    # Broad regions remain fallbacks for tilted or shifted forms and other
    # layouts. A numeric misread is not corrected by nearest-catalog guessing.
    passes.extend([
        ("item-crop", crop_code_area(grayscale), 6),
        ("po-crop", po_primary_crop(grayscale), 6),
        ("header-contrast", contrast.crop((0, 0, w, h // 2)), 11),
        ("full-page-contrast", contrast, 11),
    ])
    try:
        code, po = select_fields(texts, item_map)
        if not is_checklist and code in item_map and po:
            return code, po
        for name, region, mode in passes:
            texts[name] = pytesseract.image_to_string(
                region, lang=config.language, config=f"--psm {mode}", timeout=config.ocr_timeout)
            if name.startswith(("item-cell", "po-cell", "item-crop", "po-crop")):
                region.save(debug_dir / f"{name}.png")
            code, po = select_fields(texts, item_map)
            # Complete both template fields before accepting body/full-page text.
            if code in item_map and po and name != "item-cell-wide":
                return code, po
        for angle in (90, 180, 270):
            with contrast.rotate(angle, expand=True, fillcolor=255) as rotated:
                texts[f"rotated-{angle}"] = pytesseract.image_to_string(
                    rotated, lang=config.language, config="--psm 3", timeout=config.ocr_timeout)
                if CHECKLIST_TITLE.search(texts[f"rotated-{angle}"]):
                    for name, box, mode in CHECKLIST_CELLS:
                        key = f"{name}-rotated-{angle}"
                        with checklist_cell(rotated, box) as region:
                            texts[key] = pytesseract.image_to_string(
                                region, lang=config.language, config=f"--psm {mode}",
                                timeout=config.ocr_timeout)
                            region.save(debug_dir / f"{key}.png")
                        code, po = select_fields(texts, item_map)
                        if code in item_map and po and name != "item-cell-wide":
                            return code, po
            code, po = select_fields(texts, item_map)
            if code in item_map and po:
                return code, po
        return select_fields(texts, item_map)
    finally:
        for _, region, _ in passes:
            region.close()
        grayscale.close()
        contrast.close()


def file_signature(path):
    stat = path.stat()
    return stat.st_size, stat.st_mtime_ns


class FileChangedError(RuntimeError):
    """The scanner is still writing; leave the original in the input directory."""


def write_output(source, image, config, code, item_name, po, signature, now):
    stem = f"{now:%Y%m%d}-{code}-{sanitize_filename_component(item_name)}-{po}_"
    occurrence = 1
    while True:
        destination = config.output_folder / f"{stem}{occurrence}.pdf"
        try:
            handle = destination.open("xb")  # Exclusive creation never overwrites a scan.
            break
        except FileExistsError:
            occurrence += 1
    try:
        with handle:
            if source.suffix.lower() == ".pdf":
                with source.open("rb") as original:
                    shutil.copyfileobj(original, handle)  # Preserve ALL original PDF pages.
            else:
                image.save(handle, "PDF", resolution=config.dpi)
        if file_signature(source) != signature:
            raise FileChangedError("Scan changed during processing; waiting for it to finish")
        # The source is removed only after the complete destination was written.
        source.unlink()
        return destination
    except BaseException:
        destination.unlink(missing_ok=True)
        raise


def send_error_email(new_filename, reasons, config, debug_dir):
    password = os.environ.get("SCAN_SMTP_PASSWORD")
    if not password:
        return  # Startup and per-file logs still report problems without email configured.
    sender = os.environ.get("SCAN_EMAIL_FROM", "kc@loki-mylo.work")
    recipient = os.environ.get("SCAN_EMAIL_TO", "sconners@unlimitedsavorysystems.com")
    message = EmailMessage()
    message["From"], message["To"] = sender, recipient
    message["Subject"] = "Scan OCR needs review"
    message.set_content(
        f"File: {new_filename}\nReason: {'; '.join(reasons)}\n"
        f"Directory: {config.output_folder.resolve().as_uri()}\nDiagnostics: {debug_dir}\n")
    try:
        with smtplib.SMTP(os.environ.get("SCAN_SMTP_HOST", "smtp.zoho.com"),
                          int(os.environ.get("SCAN_SMTP_PORT", "587")), timeout=20) as server:
            server.starttls(context=ssl.create_default_context())
            server.login(os.environ.get("SCAN_SMTP_USERNAME", sender), password)
            server.send_message(message)
    except Exception as error:
        # Sorting has already completed. An email failure must never cause a retry/move.
        LOG.error("Could not email the OCR review notice for %s: %s", new_filename, error)


@dataclass(frozen=True)
class ProcessResult:
    destination: Path
    reasons: tuple[str, ...]
    debug_dir: Path


def process_file(filepath, config: Config, item_map, now=None):
    source = Path(filepath)
    if source.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported scan type: {source.suffix}")
    signature = file_signature(source)
    debug_dir = Path(tempfile.mkdtemp(
        prefix=sanitize_filename_component(source.stem)[:50] + "-", dir=config.debug_folder))
    texts = {}
    details = [f"Source: {source}", f"Workbook: {config.excel_path}",
               f"Catalog entries: {len(item_map)}"]
    image = None
    try:
        if source.suffix.lower() == ".pdf":
            try:
                texts["pdf-text"] = PdfReader(source).pages[0].extract_text() or ""
            except Exception as error:
                # Image OCR still supports PDFs with a missing/broken text layer.
                details.append(f"PDF text extraction unavailable: {error}")
        code, po = select_fields(texts, item_map)
        if not (code and po):
            image = read_first_page(source, config)
            image.save(debug_dir / "first-page.png")
            code, po = recognize_image(image, item_map, config, texts, debug_dir)
        elif source.suffix.lower() != ".pdf":
            image = read_first_page(source, config)
        reasons = []
        if not code:
            raw_codes = sorted(set().union(*(code_candidates(t)[0] for t in texts.values())))
            reasons.append("No unambiguous item number matched the Excel list" +
                           (f" (read: {', '.join(raw_codes)})" if raw_codes else ""))
        elif code not in item_map:
            reasons.append(f"Item {code} is absent from the Excel list; description is UNKNOWN")
        if not po:
            reasons.append("No unambiguous six-digit PO or STOCK found")
        code, po = code or "NO_CODE_FOUND", po or "NOPO"
        item_name = item_map.get(code, "UNKNOWN")
        if file_signature(source) != signature:
            raise FileChangedError("Scan changed during OCR; waiting for it to finish")
        destination = write_output(source, image, config, code, item_name, po,
                                   signature, now or datetime.now())
        details.extend([f"Item: {code}", f"PO: {po}", f"Output: {destination}",
                        "Review: " + ("; ".join(reasons) or "none")])
        LOG.info("Renamed and moved: %s", destination.name)
        if reasons:
            LOG.warning("OCR review needed: %s; diagnostics: %s", "; ".join(reasons), debug_dir)
            send_error_email(destination.name, reasons, config, debug_dir)
        return ProcessResult(destination, tuple(reasons), debug_dir)
    except Exception as error:
        details.append(f"Processing error: {type(error).__name__}: {error}")
        raise
    finally:
        if image is not None:
            image.close()
        try:
            with (debug_dir / "ocr.txt").open("w", encoding="utf-8") as diagnostic:
                diagnostic.write("\n".join(details) + "\n")
                for name, text in texts.items():
                    diagnostic.write(f"\n=== {name} ===\n{text}\n")
        except OSError:
            LOG.exception("Could not write diagnostics to %s", debug_dir)


def validate_setup(config):
    if config.settle_seconds < 0 or config.poll_seconds <= 0:
        raise ValueError("Settle time must be nonnegative and poll interval must be positive")
    if not config.input_folder.is_dir():
        raise ValueError(f"Input folder is unavailable: {config.input_folder}. Check the P: drive/account.")
    roots = [p.resolve() for p in (config.input_folder, config.output_folder, config.debug_folder)]
    if len(set(roots)) != 3:
        raise ValueError("Input, output, and debug folders must be different")
    for folder in (config.output_folder, config.debug_folder):
        folder.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryFile(dir=folder):
            pass  # Verify write access before starting the monitor.
    item_map = load_item_map(config.excel_path)
    pytesseract.pytesseract.tesseract_cmd = config.tesseract_path or "tesseract"
    version = pytesseract.get_tesseract_version()
    if config.language not in pytesseract.get_languages():
        raise ValueError(f"Tesseract language is not installed: {config.language}")
    for tool in ("pdfinfo", "pdftoppm"):
        if not shutil.which(tool, path=config.poppler_path):
            raise ValueError(f"Poppler {tool} is unavailable. Set SCAN_POPPLER_PATH to its bin directory.")
    LOG.info("Setup OK: Tesseract %s; item list %s (%d items); input %s",
             version, config.excel_path, len(item_map), config.input_folder)
    if not os.environ.get("SCAN_SMTP_PASSWORD"):
        LOG.warning("Email is disabled: set SCAN_SMTP_PASSWORD securely to enable review notices.")
    return item_map


@dataclass
class PendingFile:
    signature: tuple[int, int]
    stable_since: float
    retry_at: float = 0


class StabilityTracker:
    def __init__(self):
        self.pending: dict[Path, PendingFile] = {}

    def ready(self, path, now, settle_seconds):
        signature = file_signature(path)
        state = self.pending.get(path)
        if state is None or state.signature != signature:
            state = self.pending[path] = PendingFile(signature, now)
        return now - state.stable_since >= settle_seconds and now >= state.retry_at

    def defer(self, path, now, retry_seconds):
        self.pending[path].retry_at = now + retry_seconds


def incoming_files(config):
    return sorted(p for p in config.input_folder.iterdir()
                  if not p.name.startswith(".") and p.is_file() and
                  p.suffix.lower() in SUPPORTED_EXTENSIONS)


def monitor_folder(config, item_map, once=False):
    tracker = StabilityTracker()
    excel_signature = file_signature(config.excel_path)
    LOG.info("Monitoring %s; waiting %.1fs for scans to finish", config.input_folder, config.settle_seconds)
    if once:
        for source in incoming_files(config):
            tracker.ready(source, time.monotonic(), config.settle_seconds)
        time.sleep(config.settle_seconds)
    exit_code = 0
    while True:
        try:
            current_excel_signature = file_signature(config.excel_path)
            if current_excel_signature != excel_signature:
                item_map = load_item_map(config.excel_path)
                excel_signature = current_excel_signature
                LOG.info("Reloaded Excel list %s: %d items", config.excel_path, len(item_map))
            files = incoming_files(config)
            for source in list(tracker.pending):
                if source not in files:
                    del tracker.pending[source]
            for source in files:
                try:
                    if not tracker.ready(source, time.monotonic(), config.settle_seconds):
                        if once:
                            LOG.warning("Scan is still changing; left in input: %s", source.name)
                            exit_code = 1
                        continue
                    result = process_file(source, config, item_map)
                    del tracker.pending[source]
                    if result.reasons and exit_code == 0:
                        exit_code = 2
                except (OSError, FileChangedError) as error:
                    LOG.warning("Left %s in input: %s", source.name, error)
                    if source in tracker.pending:
                        tracker.defer(source, time.monotonic(), config.retry_seconds)
                    exit_code = 1
                except Exception:
                    LOG.exception("Processing failed; left %s in input", source.name)
                    tracker.defer(source, time.monotonic(), config.retry_seconds)
                    exit_code = 1
        except Exception:
            LOG.exception("Input folder or Excel list unavailable; no scans processed this cycle")
            exit_code = 1
        if once:
            return exit_code
        time.sleep(config.poll_seconds)


def main(argv=None):
    defaults = Config.from_environment()
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check", action="store_true", help="Check paths, Excel list and OCR tools without moving scans")
    parser.add_argument("--once", action="store_true", help="Process completed scans once, then exit")
    parser.add_argument("--input", type=Path, default=defaults.input_folder)
    parser.add_argument("--output", type=Path, default=defaults.output_folder)
    parser.add_argument("--debug", type=Path, default=defaults.debug_folder)
    parser.add_argument("--items", type=Path, default=defaults.excel_path)
    parser.add_argument("--tesseract", default=defaults.tesseract_path, help="Tesseract executable (otherwise PATH)")
    parser.add_argument("--poppler", default=defaults.poppler_path, help="Poppler bin directory (otherwise PATH)")
    parser.add_argument("--settle-seconds", type=float, default=defaults.settle_seconds)
    args = parser.parse_args(argv)
    config = Config(args.input, args.output, args.debug, args.items,
                    tesseract_path=args.tesseract, poppler_path=args.poppler,
                    settle_seconds=args.settle_seconds)
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    try:
        item_map = validate_setup(config)
        return 0 if args.check else monitor_folder(config, item_map, once=args.once)
    except KeyboardInterrupt:
        LOG.info("Stopped")
        return 0
    except Exception as error:
        LOG.error("Setup failed: %s", error)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
