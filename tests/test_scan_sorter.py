import importlib.util
from pathlib import Path
import sys
from datetime import datetime

import pytest
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject


SCRIPT = Path(__file__).resolve().parents[1] / "Check and Rename Incoming Scans.py"
spec = importlib.util.spec_from_file_location("scan_sorter", SCRIPT)
sorter = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = sorter
spec.loader.exec_module(sorter)

ITEMS = {"13200091": "EGG CHEESY TYPE OS NAT", "13100755": "VANILLA"}
TODAY = datetime(2026, 10, 7)


@pytest.fixture
def config(tmp_path, monkeypatch):
    monkeypatch.delenv("SCAN_SMTP_PASSWORD", raising=False)
    paths = [tmp_path / name for name in ("incoming", "sorted", "debug")]
    for path in paths:
        path.mkdir()
    workbook = Workbook()
    workbook.active.append(["Item No", "Description"])
    for code, description in ITEMS.items():
        workbook.active.append([int(code), description])
    excel_path = tmp_path / "FlavorList.xlsx"
    workbook.save(excel_path)
    workbook.close()
    return sorter.Config(*paths, excel_path=excel_path, settle_seconds=0)


@pytest.mark.parametrize("text, expected", [
    ("PO005157", "PO005157"),
    ("PO: 005157", "PO005157"),
    ("P.O. # 005157", "PO005157"),
    ("P0 00 51 57", "PO005157"),
    ("Purchase Order Number: 005157", "PO005157"),
    ("PO OO5157", "PO005157"),
    ("|PO00S157 - Int", "PO005157"),
    ("STOCK", "STOCK"),
    ("stockroom", None),
    ("Stock weight: 10 pounds", None),
    ("Purchase Order: STOCK", "STOCK"),
    ("PO #: STOCK", "STOCK"),
    ("PO: 005157\nStock", "PO005157"),
    ("PO 005157\nPO 123456", None),
    ("PO 00515789", None),
    ("PO abcdef", None),
])
def test_po_variations(text, expected):
    assert sorter.extract_po_from_text(text) == expected


@pytest.mark.parametrize("text", [
    "13200091", "Item No: 13 200 091", "Item number: 132OOO91",
    "Item No: I32OOO91", "Item # 1320-0091",
])
def test_catalog_guided_code_recognition(text):
    assert sorter.extract_code_from_text(text, ITEMS) == "13200091"


def test_codes_require_unique_matches_and_do_not_accept_dates():
    assert sorter.extract_code_from_text("13200091 13100755", ITEMS) is None
    assert sorter.extract_code_from_text("Date 20261007", ITEMS) is None
    assert sorter.extract_code_from_text("132000910", ITEMS) is None
    assert sorter.extract_code_from_text("1 3200091 0", ITEMS) is None
    assert sorter.extract_code_from_text("1 13200091", ITEMS) is None
    assert sorter.extract_code_from_text("132O0091") is None  # No unchecked substitutions.
    assert sorter.extract_code_from_text("132000\n91", ITEMS) is None


def test_explicit_item_label_wins_over_unrelated_catalog_number():
    texts = {"full-page": "Reference 13100755\nItem No: 13200091\nPO: 005157"}
    assert sorter.select_fields(texts, ITEMS) == ("13200091", "PO005157")


def test_unknown_labeled_item_keeps_number_but_unlabeled_date_does_not():
    assert sorter.select_fields({"page": "Item No: 13100746\nSTOCK"}, ITEMS) == ("13100746", "STOCK")
    assert sorter.select_fields({"page": "Date 20261007\nSTOCK"}, ITEMS) == (None, "STOCK")


def test_handwriting_reading_is_corrected_only_with_catalog_evidence():
    items = {"13100748": "CINNAMON, TOASTED WONF WS NAT"}
    assert sorter.code_candidates("32 vd J3 100748", items)[0] == {"13100748"}
    assert sorter.code_candidates("32 vd J3 100748")[0] == set()
    assert sorter.select_fields({"item-cell-wide": "32 vd J3 100748",
                                 "item-cell-tight": "14100748"}, items)[0] == "13100748"
    assert sorter.extract_code_from_text("14100748", items) is None


def test_po_cell_takes_precedence_over_body_stock_or_other_order():
    texts = {"full-page": "Item No: 13200091\nSTOCK\nExample PO123456",
             "po-cell": "|PO00S157 - Int"}
    assert sorter.select_fields(texts, ITEMS) == ("13200091", "PO005157")
    texts["po-cell"] = "STOCK"
    assert sorter.select_fields(texts, ITEMS) == ("13200091", "STOCK")


def test_checklist_body_stock_does_not_replace_an_unreadable_order():
    texts = {"full-page": "Production Checklist\nItem No: 13200091\nSTOCK\n"
                           "Stock labels printed and verified"}
    assert sorter.select_fields(texts, ITEMS) == ("13200091", None)
    texts["po-cell"] = "STOCK"
    assert sorter.select_fields(texts, ITEMS) == ("13200091", "STOCK")


def test_labeled_unknown_code_waits_for_ocr_fallback(config, monkeypatch):
    readings = iter(["Item No: 14100748\nPO005157", "Item No: 13100748\nPO005157"])
    items = {"13100748": "CINNAMON, TOASTED WONF WS NAT"}
    # Numeric conflicts are flagged, rather than accepting the first unknown
    # number or silently converting 4 to 3.
    monkeypatch.setattr(sorter.pytesseract, "image_to_string", lambda *a, **k: next(readings,
                        "Item No: 13100748\nPO005157"))
    source = config.input_folder / "scan.png"
    make_scan(source)
    result = sorter.process_file(source, config, items, TODAY)
    assert "NO_CODE_FOUND" in result.destination.name
    assert result.reasons


def test_numeric_and_text_excel_keys(config):
    workbook = Workbook()
    workbook.active.append(["Item No", "Description"])
    workbook.active.append([13200091, "EGG CHEESY TYPE OS NAT"])
    workbook.active.append(["13100755.0", "VANILLA"])
    workbook.active.append([5157, "LEADING ZERO ITEM"])
    workbook.save(config.excel_path)
    workbook.close()
    assert sorter.load_item_map(config.excel_path) == {
        **ITEMS, "00005157": "LEADING ZERO ITEM"}


def test_bad_catalog_fails_instead_of_silent_unknowns(config):
    with pytest.raises(FileNotFoundError):
        sorter.load_item_map(config.excel_path.with_name("missing.xlsx"))
    workbook = Workbook()
    workbook.active.append(["Item No", "Description"])
    workbook.active.append([13200091, "First"])
    workbook.active.append([13200091, "Conflicting"])
    workbook.save(config.excel_path)
    workbook.close()
    with pytest.raises(ValueError, match="conflicting"):
        sorter.load_item_map(config.excel_path)


def make_scan(path, *, faint=False, angle=0, pages=1):
    image = Image.new("RGB", (1650, 2200), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype("DejaVuSans.ttf", 44)
    color = (170, 170, 170) if faint else "black"
    draw.text((120, 100), "CHECKLIST", fill=color, font=font)
    # Both fields are outside the original fixed crops.
    draw.text((160, 1100), "Item No: 13200091", fill=color, font=font)
    draw.text((800, 1700), "P.O. # 005157", fill=color, font=font)
    if angle:
        rotated = image.rotate(angle, expand=True, fillcolor="white")
        image.close()
        image = rotated
    second = Image.new("RGB", image.size, "white")
    try:
        if path.suffix == ".pdf":
            image.save(path, "PDF", resolution=200, save_all=pages > 1,
                       append_images=[second] if pages > 1 else [])
        else:
            image.save(path)
    finally:
        image.close()
        second.close()


def fake_recognition(monkeypatch, texts="Item No: 13200091\nPO: 005157"):
    monkeypatch.setattr(sorter.pytesseract, "image_to_string", lambda *a, **kw: texts)


def test_pdf_keeps_all_pages_and_duplicate_filenames(config, monkeypatch):
    fake_recognition(monkeypatch)
    source = config.input_folder / "scan.pdf"
    make_scan(source, pages=2)
    original = source.read_bytes()
    result = sorter.process_file(source, config, ITEMS, TODAY)
    assert not source.exists()
    assert not result.reasons
    assert result.destination.read_bytes() == original
    assert len(PdfReader(result.destination).pages) == 2
    source.write_bytes(original)
    second = sorter.process_file(source, config, ITEMS, TODAY)
    assert second.destination.name.endswith("_2.pdf")
    assert result.destination.read_bytes() == original
    assert second.destination.read_bytes() == original


def test_searchable_pdf_uses_existing_text_without_rasterization(config, monkeypatch):
    source = config.input_folder / "searchable.pdf"
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                             NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"):
        DictionaryObject({NameObject("/F1"): font})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 14 Tf 50 700 Td (Item No: 13200091) Tj "
                    b"0 -30 Td (PO: 005157) Tj ET")
    page[NameObject("/Contents")] = stream
    writer.write(source)
    def unexpected_render(*args, **kwargs):
        pytest.fail("Searchable PDF should not need rendering/OCR")
    monkeypatch.setattr(sorter, "read_first_page", unexpected_render)
    original = source.read_bytes()
    result = sorter.process_file(source, config, ITEMS, TODAY)
    assert not result.reasons
    assert result.destination.read_bytes() == original


def test_png_is_converted_to_real_pdf(config, monkeypatch):
    fake_recognition(monkeypatch)
    source = config.input_folder / "scan.png"
    make_scan(source)
    result = sorter.process_file(source, config, ITEMS, TODAY)
    assert result.destination.read_bytes().startswith(b"%PDF-")
    assert len(PdfReader(result.destination).pages) == 1
    assert not source.exists()


def test_unknown_item_is_reported_separately_from_missing_ocr(config, monkeypatch):
    fake_recognition(monkeypatch, "Item No: 13100746\nPO: 005157")
    notices = []
    monkeypatch.setattr(sorter, "send_error_email", lambda *args: notices.append(args))
    source = config.input_folder / "scan.png"
    make_scan(source)
    result = sorter.process_file(source, config, ITEMS, TODAY)
    assert "13100746-UNKNOWN-PO005157" in result.destination.name
    assert result.reasons == ("Item 13100746 is absent from the Excel list; description is UNKNOWN",)
    assert len(notices) == 1


def test_missing_fields_generate_one_explanatory_notice(config, monkeypatch):
    fake_recognition(monkeypatch, "unreadable checklist")
    notices = []
    monkeypatch.setattr(sorter, "send_error_email", lambda *args: notices.append(args))
    source = config.input_folder / "scan.png"
    make_scan(source)
    result = sorter.process_file(source, config, ITEMS, TODAY)
    assert "NO_CODE_FOUND-UNKNOWN-NOPO" in result.destination.name
    assert len(result.reasons) == 2
    assert len(notices) == 1
    assert "rotated-270" in (result.debug_dir / "ocr.txt").read_text()
    assert "Review:" in (result.debug_dir / "ocr.txt").read_text()


def test_ocr_failure_leaves_input_and_sends_no_email(config, monkeypatch):
    source = config.input_folder / "scan.png"
    make_scan(source)
    original = source.read_bytes()
    notices = []
    def failure(*a, **kw):
        raise RuntimeError("Tesseract timeout")
    monkeypatch.setattr(sorter.pytesseract, "image_to_string", failure)
    monkeypatch.setattr(sorter, "send_error_email", lambda *args: notices.append(args))
    with pytest.raises(RuntimeError, match="timeout"):
        sorter.process_file(source, config, ITEMS, TODAY)
    assert source.read_bytes() == original
    assert not list(config.output_folder.iterdir())
    assert not notices
    assert "Tesseract timeout" in next(config.debug_folder.glob("*/ocr.txt")).read_text()


def test_changing_file_waits_and_resets_settle_timer(config):
    source = config.input_folder / "scan.pdf"
    source.write_bytes(b"first chunk")
    tracker = sorter.StabilityTracker()
    assert not tracker.ready(source, 0, 10)
    assert not tracker.ready(source, 5, 10)
    source.write_bytes(b"first chunk plus second chunk")
    assert not tracker.ready(source, 9, 10)
    assert not tracker.ready(source, 18, 10)
    assert tracker.ready(source, 19, 10)
    tracker.defer(source, 19, 60)
    assert not tracker.ready(source, 78, 10)
    assert tracker.ready(source, 79, 10)


def test_source_changes_during_copy_roll_back_destination(config, monkeypatch):
    source = config.input_folder / "scan.pdf"
    source.write_bytes(b"%PDF-original")
    signature = sorter.file_signature(source)
    def changed_copy(original, destination):
        destination.write(original.read())
        source.write_bytes(b"%PDF-new scanner data")
    monkeypatch.setattr(sorter.shutil, "copyfileobj", changed_copy)
    with pytest.raises(sorter.FileChangedError):
        sorter.write_output(source, None, config, "13200091", ITEMS["13200091"],
                            "PO005157", signature, TODAY)
    assert source.read_bytes() == b"%PDF-new scanner data"
    assert not list(config.output_folder.iterdir())


def test_locked_source_leaves_original_and_rolls_back_output(config, monkeypatch):
    source = config.input_folder / "scan.pdf"
    source.write_bytes(b"%PDF-original")
    signature = sorter.file_signature(source)
    original_unlink = Path.unlink
    def locked_unlink(path, *args, **kwargs):
        if path == source:
            raise PermissionError("File is still open in the scanner")
        return original_unlink(path, *args, **kwargs)
    monkeypatch.setattr(Path, "unlink", locked_unlink)
    with pytest.raises(PermissionError, match="still open"):
        sorter.write_output(source, None, config, "13200091", ITEMS["13200091"],
                            "PO005157", signature, TODAY)
    assert source.read_bytes() == b"%PDF-original"
    assert not list(config.output_folder.iterdir())


@pytest.mark.parametrize("suffix, faint, angle", [
    (".png", False, 0), (".pdf", False, 0), (".png", True, 0),
    (".png", False, 90), (".png", False, 180),
])
def test_real_ocr_changed_layout_and_scan_quality(config, suffix, faint, angle):
    sorter.validate_setup(config)
    source = config.input_folder / f"scan{suffix}"
    make_scan(source, faint=faint, angle=angle)
    result = sorter.process_file(source, config, ITEMS, TODAY)
    assert result.destination.name == "20261007-13200091-EGG CHEESY TYPE OS NAT-PO005157_1.pdf"
    assert not result.reasons
    assert not source.exists()


def test_once_reports_review_exit_status(config, monkeypatch):
    fake_recognition(monkeypatch, "Item No: 13200091")
    source = config.input_folder / "scan.png"
    make_scan(source)
    assert sorter.monitor_folder(config, ITEMS, once=True) == 2
    assert not source.exists()
